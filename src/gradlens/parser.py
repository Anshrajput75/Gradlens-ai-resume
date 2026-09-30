"""Local document ingestion and transparent, rule-based resume structuring."""

from __future__ import annotations

import re
from pathlib import Path

from .models import ResumeProfile

SECTION_ALIASES = {
    "summary": ("summary", "profile", "objective", "research interests"),
    "education": ("education", "academic background"),
    "experience": ("experience", "employment", "work history", "professional experience"),
    "research": ("research", "research experience", "laboratory experience"),
    "projects": ("projects", "selected projects", "academic projects"),
    "publications": ("publications", "papers", "presentations"),
    "skills": ("skills", "technical skills", "competencies"),
    "awards": ("awards", "honors", "honours", "scholarships"),
    "leadership": ("leadership", "service", "volunteering", "activities"),
    "certifications": ("certifications", "certificates", "training"),
}
_HEADER_MAP = {alias: section for section, aliases in SECTION_ALIASES.items() for alias in aliases}
_EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
_PHONE = re.compile(r"(?<!\w)(?:\+?\d[\d ().-]{7,}\d)(?!\w)")
_DATE = re.compile(
    r"\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?\.?\s+)?"
    r"(?:19|20)\d{2}(?:\s*[-–—]\s*(?:(?:19|20)\d{2}|Present|Current))?\b",
    re.I,
)
_SKILL_SPLIT = re.compile(r"[,;|•]")
_BULLET = re.compile(r"^\s*(?:[•●▪◦*\-]|\d+[.)])\s*")


def extract_text(path: Path) -> str:
    """Extract UTF-8 text, or use the optional reader matching the file extension."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support is optional. Install it with: pip install 'gradlens[pdf]'") from exc
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise ValueError("No selectable text found. Scanned PDFs need OCR before parsing.")
        return text
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support is optional. Install it with: pip install 'gradlens[docx]'") from exc
        document = Document(str(path))
        chunks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            chunks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(chunks)
    raise ValueError(f"Unsupported file type '{suffix}'. Use .txt, .pdf, or .docx.")


def _normalize_header(line: str) -> str | None:
    candidate = re.sub(r"[:\s]+$", "", line.strip().lower())
    candidate = re.sub(r"\s+", " ", candidate)
    return _HEADER_MAP.get(candidate)


def parse_text(text: str, source_file: str | None = None) -> ResumeProfile:
    """Structure resume text into common sections; extraction is heuristic."""
    lines = [line.strip() for line in text.replace("\r", "\n").splitlines()]
    sections: dict[str, list[str]] = {}
    current = "contact"
    sections[current] = []
    for line in lines:
        if not line:
            continue
        header = _normalize_header(line)
        if header:
            current = header
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(_BULLET.sub("", line).strip())

    email_match = _EMAIL.search(text)
    phone_match = _PHONE.search(text)
    contact_lines = sections.get("contact", [])
    name = next(
        (line for line in contact_lines if not _EMAIL.search(line) and not _PHONE.search(line) and len(line) < 90),
        None,
    )
    location = None
    for line in contact_lines:
        remainder = _EMAIL.sub("", line)
        remainder = _PHONE.sub("", remainder)
        remainder = re.sub(r"^[\s|,;•-]+|[\s|,;•-]+$", "", remainder).strip()
        if remainder and remainder != name:
            location = remainder
            break
    skills = []
    for line in sections.get("skills", []):
        skills.extend(part.strip() for part in _SKILL_SPLIT.split(line) if part.strip())
    skills = list(dict.fromkeys(skills))
    dates = list(dict.fromkeys(_DATE.findall(text)))
    warnings = []
    if not sections.get("education"):
        warnings.append("No Education section was detected; check whether the heading uses a supported label.")
    if not sections.get("experience") and not sections.get("research"):
        warnings.append("No Experience or Research section was detected; verify section headings.")
    if not name:
        warnings.append("Name could not be confidently extracted; review the parsed profile.")
    return ResumeProfile(
        name=name,
        email=email_match.group(0) if email_match else None,
        phone=phone_match.group(0).strip() if phone_match else None,
        location=location,
        sections={k: v for k, v in sections.items() if k != "contact"},
        skills=skills,
        dates=dates,
        word_count=len(re.findall(r"\b[\w'-]+\b", text)),
        source_file=source_file,
        warnings=warnings,
    )


def parse_file(path: str | Path) -> ResumeProfile:
    file_path = Path(path).expanduser()
    if not file_path.is_file():
        raise FileNotFoundError(f"Resume file not found: {file_path}")
    return parse_text(extract_text(file_path), source_file=file_path.name)
